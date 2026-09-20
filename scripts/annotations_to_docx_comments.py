#!/usr/bin/env python3
"""Move the final internal annotation chapter of a DOCX into real Word comments."""

from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsmap, qn
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.text.paragraph import Paragraph


W = nsmap["w"]
COMMENTS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
COMMENTS_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
ANNOTATION_TITLES = {"批注", "批注（判断过程、思路与假设）"}
MARKDOWN_SUBSECTION_RE = re.compile(r"^###\s+\d+[.、][ \t]*(.+?)[ \t]*$", re.MULTILINE)


def _paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _heading_level(element) -> int | None:
    if element.tag != qn("w:p"):
        return None
    properties = element.find(qn("w:pPr"))
    if properties is None:
        return None
    style = properties.find(qn("w:pStyle"))
    if style is None:
        return None
    value = style.get(qn("w:val"), "")
    match = re.fullmatch(r"Heading ?([1-9])", value, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _annotation_start(children) -> tuple[int, int] | None:
    """Find only a heading whose text is the annotation chapter title."""
    for index, element in enumerate(children):
        level = _heading_level(element)
        if level is None:
            continue
        text = _paragraph_text(element).strip()
        if text in ANNOTATION_TITLES:
            return index, level
    return None


def _element_text(element) -> str:
    if element.tag == qn("w:tbl"):
        row_texts: list[str] = []
        for row in element.findall(qn("w:tr")):
            cells = []
            for cell in row.findall(qn("w:tc")):
                cell_text = "\n".join(
                    _paragraph_text(paragraph).strip()
                    for paragraph in cell.findall(qn("w:p"))
                    if _paragraph_text(paragraph).strip()
                )
                cells.append(cell_text)
            if any(cells):
                row_texts.append(" | ".join(cells))
        return "\n".join(row_texts)
    return _paragraph_text(element)


def _join_elements(elements) -> str:
    parts = [_element_text(element).strip() for element in elements]
    return "\n".join(part for part in parts if part)


def parse_annotations(markdown_text: str) -> list[tuple[str, str]]:
    """Parse the Markdown representation of an annotation chapter for callers that need it."""
    matches = list(MARKDOWN_SUBSECTION_RE.finditer(markdown_text))
    if not matches:
        body = markdown_text.strip()
        return [("批注", body)] if body else []
    entries: list[tuple[str, str]] = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown_text)
        body = markdown_text[match.end() : end].strip()
        if body:
            entries.append((match.group(1).strip(), body))
    return entries


def extract_and_strip_annotations(doc: Document) -> list[tuple[str, str]]:
    """Extract a structured final annotation chapter and leave section properties intact."""
    body = doc.element.body
    children = list(body)
    found = _annotation_start(children)
    if found is None:
        return []
    start, container_level = found
    end = len(children)
    for index in range(start, len(children)):
        if children[index].tag == qn("w:sectPr"):
            end = index
            break

    removed = children[start:end]
    entries: list[tuple[str, str]] = []
    current_title: str | None = None
    current_body: list = []
    container_body: list = []

    for element in removed[1:]:
        level = _heading_level(element)
        if level is not None and level > container_level:
            if current_title is not None:
                body_text = _join_elements(current_body)
                if body_text:
                    entries.append((current_title, body_text))
            current_title = _paragraph_text(element).strip()
            current_body = []
        elif current_title is None:
            container_body.append(element)
        else:
            current_body.append(element)

    if current_title is not None:
        body_text = _join_elements(current_body)
        if body_text:
            entries.append((current_title, body_text))
    elif container_body:
        body_text = _join_elements(container_body)
        if body_text:
            entries.append(("批注", body_text))

    for element in removed:
        body.remove(element)
    return entries


def _get_or_create_comments_part(doc: Document):
    for relationship in doc.part.rels.values():
        if relationship.reltype == COMMENTS_REL:
            return relationship.target_part
    comments_element = parse_xml(f'<w:comments xmlns:w="{W}"/>')
    comments_part = XmlPart(
        PackURI("/word/comments.xml"),
        COMMENTS_CONTENT_TYPE,
        comments_element,
        doc.part.package,
    )
    doc.part.relate_to(comments_part, COMMENTS_REL)
    return comments_part


def _ensure_comment_styles(doc: Document) -> None:
    styles = doc.styles.element
    present = {style.get(qn("w:styleId")) for style in styles.findall(qn("w:style"))}
    for style_id, style_type, display_name in (
        ("CommentReference", "character", "Comment Reference"),
        ("CommentText", "paragraph", "Comment Text"),
    ):
        if style_id in present:
            continue
        style = OxmlElement("w:style")
        style.set(qn("w:type"), style_type)
        style.set(qn("w:styleId"), style_id)
        name = OxmlElement("w:name")
        name.set(qn("w:val"), display_name)
        style.append(name)
        styles.append(style)


def _next_comment_id(comments_part) -> int:
    ids = []
    for comment in comments_part.element.iter(qn("w:comment")):
        raw_id = comment.get(qn("w:id"))
        if raw_id is not None and raw_id.lstrip("-").isdigit():
            ids.append(int(raw_id))
    return max(ids, default=-1) + 1


def _paragraphs(doc: Document) -> list[Paragraph]:
    return [Paragraph(element, doc) for element in doc.element.body.findall(qn("w:p"))]


def _find_anchor(doc: Document, title: str) -> Paragraph:
    """定位批注锚点段落。

    H2 质量加固（2026-09-15）：原实现在无法定位时静默兜底（挂末段或新建空段），
    批注可能挂错位置且无任何提示。现改为在兜底分支写 stderr 告警，仍保持兜底
    行为（不中断转换）。
    """
    # 锚点规则须与 assets/modeA-report-template.md / assets/modeB-comparison-template.md
    # 的「正文章节名」及「批注子节名」对齐（2026-09-18 修订：旧规则引用的
    # 「已知与待核事实/来源与效力核验/适用性结果/产品落地清单」系旧版模板章节名，
    # 当前模板正文已改为「一、画像摘要与角色判定」/「六、效力核验记录」/
    # 「三、分法域义务详述」，导致四个批注子节全部锚定失败降级挂末段）。
    # 模式A 批注四子节 → 正文锚点：
    #   ① 角色判定推理（双锚定）→ 一、画像摘要与角色判定
    #   ② 前置判断与适用性结论 → 一、画像摘要与角色判定（适用性结论归属画像节）
    #   ③ 假设与不确定项        → 六、效力核验记录（待核事实/假设留痕）
    #   ④ 特定规则适用说明      → 三、分法域义务详述（逐法域规则适用）
    # 模式B 批注子节 → 正文锚点（2026-09-19 补：modeB 对比报告原无映射，
    #   批注全部降级挂末段；按 modeB-comparison-template 章节名补映射）：
    #   ① 概念对齐与可比性说明   → 二、概念对齐
    #   ② 主体对应关系推理       → 四、义务主体对应关系
    #   ③ 假设与不确定项         → 九、效力核验记录（modeB 效力节序号不同）
    #   ④ 特定规则适用说明       → 三、对比矩阵（逐维规则适用）
    #   ⑤ 行为准则适用与对比维度说明 → 五、企业合规义务梳理（CoP 措施就地展开处；2026-09-19
    #      用户裁定取消独立 CoP 章节后，由「八、行为准则措施要点展开」改锚此处）
    #      单法域梳理模板无「企业合规义务梳理」节，其 CoP 措施落在各主体「成文法义务与行为准则义务的衔接」
    #      小节内，故该规则同时保留后者为锚点词（两模板共用一条规则，按文档实际存在的段落命中）。
    # 模式B 单法域梳理模板（2026-09-20 补：该模板批注 ①③ 原无锚点规则，实测两条降级挂文档末段；
    #   用户 2026-09-20 裁定两条均锚「二、义务主体与义务内容」）：
    #   ① 主体判定推理       → 二、义务主体与义务内容
    #   ③ 特定规则适用说明   → 二、义务主体与义务内容
    #   注：① 的标题词「主体」「判定」与 modeA ①（角色判定推理，经「判定」命中）及
    #   对比模板 ②（主体对应关系推理，经「主体」命中）部分重叠；因规则**按序尝试、
    #   不提前放弃**，且两条新规则置于末位，故 modeA/对比模板的既有命中不受影响。
    # 规则按序尝试：先命中的标题词若无对应锚点段落，则继续尝试后续规则（不提前放弃）。
    title_rules = (
        (("角色", "判定"), ("画像摘要与角色判定",)),
        (("前置", "判断"), ("画像摘要与角色判定",)),
        (("假设", "不确定"), ("效力核验记录",)),
        (("规则", "说明"), ("分法域义务详述",)),
        (("概念", "对齐"), ("概念对齐",)),
        (("主体", "对应"), ("义务主体对应关系",)),
        (("对比维度", "准则"), ("企业合规义务梳理", "成文法义务与行为准则义务的衔接")),
        (("规则", "说明"), ("对比矩阵",)),
        # 模式B 单法域梳理模板（2026-09-20 补；置于末位，不影响上方 modeA／对比模板命中）
        (("主体", "判定"), ("义务主体与义务内容",)),
        (("特定规则", "适用"), ("义务主体与义务内容",)),
    )
    paragraphs = _paragraphs(doc)
    matched_rule = False
    for title_words, anchor_words in title_rules:
        if any(word in title for word in title_words):
            matched_rule = True
            for paragraph in paragraphs:
                if any(word in paragraph.text for word in anchor_words):
                    return paragraph
    if paragraphs:
        reason = "未匹配到锚点段落" if matched_rule else "标题不属于已知归类"
        print(
            f"warning: 批注锚定失败（{reason}），已降级挂到文档末段：{title}",
            file=sys.stderr,
        )
        return paragraphs[-1]
    print(f"warning: 文档无可定位段落，已新建空段落承载批注：{title}", file=sys.stderr)
    return doc.add_paragraph()


def _append_comment_text(comment, text: str) -> None:
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), "CommentText")
    properties.append(style)
    paragraph.append(properties)
    run = OxmlElement("w:r")
    for index, line in enumerate(text.split("\n")):
        if line:
            node = OxmlElement("w:t")
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            node.text = line
            run.append(node)
        if index < len(text.split("\n")) - 1:
            run.append(OxmlElement("w:br"))
    if len(run) == 0:
        run.append(OxmlElement("w:t"))
    paragraph.append(run)
    comment.append(paragraph)


def _add_comment(comments_part, anchor: Paragraph, comment_id: int, title: str, body: str, author: str) -> None:
    if not anchor._p.findall(qn("w:r")):
        anchor.add_run(" ")
    start = OxmlElement("w:commentRangeStart")
    start.set(qn("w:id"), str(comment_id))
    # OOXML CT_P 规定 w:pPr 必须是段落的首个子元素，其后才是内容组；故锚点须插到
    # w:pPr 之后（2026-09-19 修复 S-2：原 `insert(0, start)` 落在 w:pPr 之前，
    # 违反 schema，严格校验器/部分 Word 版本会提示「内容有问题，需要修复」）。
    anchor_p = anchor._p
    insert_at = 1 if anchor_p.find(qn("w:pPr")) is not None else 0
    anchor_p.insert(insert_at, start)
    end = OxmlElement("w:commentRangeEnd")
    end.set(qn("w:id"), str(comment_id))
    reference_run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "CommentReference")
    properties.append(style)
    reference_run.append(properties)
    reference = OxmlElement("w:commentReference")
    reference.set(qn("w:id"), str(comment_id))
    reference_run.append(reference)
    anchor._p.append(end)
    anchor._p.append(reference_run)

    comment = OxmlElement("w:comment")
    comment.set(qn("w:id"), str(comment_id))
    comment.set(qn("w:author"), author)
    comment.set(qn("w:initials"), author[:4] or "AI")
    comment.set(qn("w:date"), datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"))
    _append_comment_text(comment, f"{title}\n\n{body}")
    comments_part.element.append(comment)


def convert_docx_annotations(docx_path: str | Path, author: str = "合规分析") -> int:
    """Move one final annotation chapter into comments and return the new count."""
    path = Path(docx_path)
    doc = Document(path)
    had_annotation_chapter = _annotation_start(list(doc.element.body)) is not None
    entries = extract_and_strip_annotations(doc)
    if not had_annotation_chapter:
        return 0
    if not entries:
        doc.save(path)
        return 0
    _ensure_comment_styles(doc)
    comments_part = _get_or_create_comments_part(doc)
    comment_id = _next_comment_id(comments_part)
    added = 0
    for title, body in entries:
        if not body.strip():
            continue
        _add_comment(comments_part, _find_anchor(doc, title), comment_id, title, body, author)
        comment_id += 1
        added += 1
    doc.save(path)
    return added


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="将末尾批注章节转换为 Word 批注")
    parser.add_argument("docx", help="待处理的 DOCX 文件")
    parser.add_argument("--author", default="合规分析", help="批注作者")
    args = parser.parse_args(argv)
    count = convert_docx_annotations(args.docx, author=args.author)
    print(f"converted {count} annotation(s)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
