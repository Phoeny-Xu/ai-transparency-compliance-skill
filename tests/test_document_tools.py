from __future__ import annotations

import importlib.util
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree


SKILL_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = SKILL_ROOT / "scripts"
PYTHON = sys.executable


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


MD_TO_DOCX = load_module("skill_md_to_docx", SCRIPTS / "md_to_docx.py")
ANNOTATIONS = load_module("skill_annotations", SCRIPTS / "annotations_to_docx_comments.py")
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


class DocumentToolTests(unittest.TestCase):
    def convert(self, markdown: str, *, annotations: bool = True) -> Path:
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        source = Path(directory.name) / "report.md"
        target = Path(directory.name) / "report.docx"
        source.write_text(markdown, encoding="utf-8")
        MD_TO_DOCX.convert(source, target, convert_annotations=annotations)
        return target

    def package_xml(self, path: Path, member: str) -> etree._Element:
        with zipfile.ZipFile(path) as archive:
            return etree.fromstring(archive.read(member))

    def package_names(self, path: Path) -> set[str]:
        with zipfile.ZipFile(path) as archive:
            return set(archive.namelist())

    def test_table_followers_and_table_variants_are_preserved(self):
        path = self.convert(
            "| A | B |\n| --- | --- |\n| left | right |\n"
            "## 表格后的标题\n表格后的段落\n- 表格后的列表\n\n"
            "A | B\n--- | ---\na\\|b | `x|y` | 多出的单元格\n"
        )
        document = Document(path)
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("表格后的标题", paragraphs)
        self.assertIn("表格后的段落", paragraphs)
        self.assertIn("表格后的列表", paragraphs)
        self.assertEqual(len(document.tables), 2)
        self.assertEqual(len(document.tables[1].columns), 3)
        self.assertEqual(document.tables[1].cell(1, 0).text, "a|b")
        self.assertEqual(document.tables[1].cell(1, 1).text, "x|y")

    def test_footnotes_support_continuations_formatting_and_orphan_filtering(self):
        path = self.convert(
            "正文[^used]\n\n"
            "[^used]: 第一行 **加粗**\n"
            "    第二行\n"
            "[^orphan]: 不应写入 Word 包\n"
        )
        with zipfile.ZipFile(path) as archive:
            footnotes = archive.read("word/footnotes.xml")
        self.assertIn("第一行", footnotes.decode("utf-8"))
        self.assertIn("第二行", footnotes.decode("utf-8"))
        self.assertNotIn("不应写入", footnotes.decode("utf-8"))
        root = etree.fromstring(footnotes)
        self.assertTrue(root.findall(".//{%s}b" % W_NS))

    def test_undefined_or_duplicate_footnotes_fail_loudly(self):
        with self.assertRaisesRegex(ValueError, "Undefined footnote reference"):
            self.convert("正文[^missing]\n")
        with self.assertRaisesRegex(ValueError, "Duplicate footnote definition"):
            self.convert("正文[^same]\n[^same]: one\n[^same]: two\n")

    def test_heading_inline_code_and_six_levels(self):
        path = self.convert(
            "# **粗体** [链接](https://example.com)[^h]\n"
            "##### 五级\n###### 六级\n"
            "正文含 `inline()`。\n\n```python\nprint('fenced')\n```\n"
            "[^h]: heading note\n"
        )
        document = Document(path)
        headings = [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs]
        self.assertIn(("Heading 5", "五级"), headings)
        self.assertIn(("Heading 6", "六级"), headings)
        all_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertNotIn("`", all_text)
        self.assertIn("inline()", all_text)
        self.assertIn("print('fenced')", all_text)
        with zipfile.ZipFile(path) as archive:
            document_xml = archive.read("word/document.xml")
        self.assertIn(b"w:hyperlink", document_xml)
        self.assertIn(b"footnoteReference", document_xml)

    def test_links_keep_balanced_urls_and_support_relative_and_mailto(self):
        path = self.convert(
            "裸链接 https://example.com/path_(one).\n"
            "[相对](../policy) [邮件](mailto:legal@example.com)\n"
        )
        with zipfile.ZipFile(path) as archive:
            rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        targets = {relationship.get("Target") for relationship in rels}
        self.assertIn("https://example.com/path_(one)", targets)
        self.assertNotIn("https://example.com/path_(one).", targets)
        self.assertIn("../policy", targets)
        self.assertIn("mailto:legal@example.com", targets)

    def test_invalid_color_never_crashes_conversion(self):
        path = self.convert('<font color="#x">颜色</font><span style="color:#12">也保留</span>')
        self.assertIn("颜色也保留", "".join(paragraph.text for paragraph in Document(path).paragraphs))

    def test_cli_requires_both_paths(self):
        result = subprocess.run(
            [PYTHON, str(SCRIPTS / "md_to_docx.py")],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 2)
        self.assertIn("usage", result.stderr.lower())

    def test_default_annotation_hook_moves_only_the_final_heading_chapter(self):
        path = self.convert(
            "普通正文含有 批注（判断过程、思路与假设） 这几个字，但不是标题。\n\n"
            "## 结论\n内容\n\n"
            "## 批注（判断过程、思路与假设）\n\n"
            "### 1. 事实说明\n第一行\n第二行\n"
        )
        names = self.package_names(path)
        self.assertIn("word/comments.xml", names)
        document_xml = self.package_xml(path, "word/document.xml")
        self.assertTrue(document_xml.findall(".//{%s}sectPr" % W_NS))
        document_text = "".join(document_xml.itertext())
        self.assertIn("普通正文含有", document_text)
        self.assertNotIn("事实说明", document_text)
        comments = self.package_xml(path, "word/comments.xml")
        comment_text = "".join(comments.itertext())
        self.assertIn("第一行", comment_text)
        self.assertIn("第二行", comment_text)
        self.assertTrue(comments.findall(".//{%s}br" % W_NS))

    def test_empty_annotation_chapter_creates_no_fake_comment(self):
        path = self.convert("正文\n\n## 批注（判断过程、思路与假设）\n\n### 1. 空白\n\n### 2. 仍为空白\n")
        self.assertNotIn("word/comments.xml", self.package_names(path))
        self.assertNotIn("批注", "".join(paragraph.text for paragraph in Document(path).paragraphs))

    def test_existing_comment_ids_remain_unique_and_author_is_xml_safe(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "comments.docx"
        document = Document()
        document.add_heading("正文", level=1)
        document.add_heading("批注（判断过程、思路与假设）", level=2)
        document.add_heading("1. 第一条", level=3)
        document.add_paragraph("first")
        document.save(path)
        self.assertEqual(ANNOTATIONS.convert_docx_annotations(path, author='A&B "<x>'), 1)

        document = Document(path)
        document.add_heading("批注（判断过程、思路与假设）", level=2)
        document.add_heading("1. 第二条", level=3)
        document.add_paragraph("second")
        document.save(path)
        self.assertEqual(ANNOTATIONS.convert_docx_annotations(path), 1)

        comments = self.package_xml(path, "word/comments.xml")
        comments_nodes = comments.findall(".//{%s}comment" % W_NS)
        ids = [node.get("{%s}id" % W_NS) for node in comments_nodes]
        self.assertEqual(len(ids), len(set(ids)))
        self.assertEqual(comments_nodes[0].get("{%s}author" % W_NS), 'A&B "<x>')

    def test_comments_have_a_fallback_anchor_when_no_body_paragraph_exists(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        path = Path(directory.name) / "table-only.docx"
        document = Document()
        document.add_table(rows=1, cols=1).cell(0, 0).text = "only table"
        document.add_heading("批注（判断过程、思路与假设）", level=2)
        document.add_heading("1. 表格说明", level=3)
        document.add_paragraph("说明")
        document.save(path)
        self.assertEqual(ANNOTATIONS.convert_docx_annotations(path), 1)
        self.assertIn("word/comments.xml", self.package_names(path))
        self.assertTrue(self.package_xml(path, "word/document.xml").findall(".//{%s}commentReference" % W_NS))


if __name__ == "__main__":
    unittest.main()
